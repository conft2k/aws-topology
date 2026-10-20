#!/usr/bin/env python3
"""Render each slide of a deck made by svg2pptx.py to PNG, for a visual check.

    uv run --with python-pptx --with cairosvg python scripts/render_pptx.py deck.pptx preview

Writes preview-1.png, preview-2.png, ... (one per slide). It reads the shapes back out
of the .pptx — positions, fills, borders, dashes, arrowheads, text boxes and the SVG
icons — so it shows what the file actually contains, not the source SVG. It is not
PowerPoint: text is drawn with the first installed Korean font (Noto Sans CJK KR /
Pretendard) at the box position PowerPoint uses for a single line, so expect small
differences in text width against the real typeface. Use it to catch conversion
problems; open the deck in PowerPoint or Keynote for the final look.
"""
import base64
import subprocess
import sys
from pathlib import Path
from xml.sax.saxutils import escape

import cairosvg
from pptx import Presentation
from pptx.oxml.ns import qn

ASVG_BLIP = "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip"
REL_EMBED = qn("r:embed")
ASCENT = 0.95  # keep in sync with svg2pptx.py


def preview_font():
    try:
        out = subprocess.run(["fc-list"], capture_output=True, text=True).stdout
    except OSError:
        out = ""
    for fam in ("Noto Sans CJK KR", "Pretendard"):
        if fam in out:
            return fam
    return "sans-serif"


class Renderer:
    def __init__(self, slide, font):
        self.slide, self.font, self.out = slide, font, []
        self.markers = set()

    # group transform: child coords -> slide coords
    def walk(self, shapes, tf):
        for shape in shapes:
            el = shape._element
            tag = el.tag.split("}")[1]
            if tag == "grpSp":
                xfrm = el.grpSpPr.find(qn("a:xfrm"))
                off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
                choff, chext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
                sx = int(ext.get("cx")) / max(1, int(chext.get("cx")))
                sy = int(ext.get("cy")) / max(1, int(chext.get("cy")))
                ox, oy = int(off.get("x")), int(off.get("y"))
                cx, cy = int(choff.get("x")), int(choff.get("y"))
                inner = (lambda f, ox=ox, oy=oy, cx=cx, cy=cy, sx=sx, sy=sy:
                         lambda x, y, w, h: f(ox + (x - cx) * sx, oy + (y - cy) * sy, w * sx, h * sy))(tf)
                self.walk(shape.shapes, inner)
            elif tag == "pic":
                self.pic(el, tf)
            elif tag in ("sp", "cxnSp"):
                self.sp(el, tf, tag == "cxnSp")

    @staticmethod
    def box(el, tf):
        xfrm = el.find(".//" + qn("a:xfrm"))
        off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
        x, y, w, h = tf(int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))
        return x, y, w, h, xfrm.get("flipH") == "1", xfrm.get("flipV") == "1"

    def stroke(self, sp_pr):
        ln = sp_pr.find(qn("a:ln"))
        if ln is None or ln.find(qn("a:noFill")) is not None or ln.find(qn("a:solidFill")) is None:
            return 'stroke="none"'
        w = int(ln.get("w", "12700"))
        c = ln.find(qn("a:solidFill")).find(qn("a:srgbClr")).get("val")
        attrs = f'stroke="#{c}" stroke-width="{w}"'
        cust = ln.find(qn("a:custDash"))
        if cust is not None:
            dash = []
            for ds in cust.findall(qn("a:ds")):
                dash += [int(ds.get("d")) * w / 100000, int(ds.get("sp")) * w / 100000]
            attrs += f' stroke-dasharray="{",".join(f"{v:.0f}" for v in dash)}"'
        for end, which in (("a:headEnd", "start"), ("a:tailEnd", "end")):
            e = ln.find(qn(end))
            if e is not None and e.get("type", "none") != "none":
                mid = f"arr-{c}-{which}"
                self.markers.add((mid, c, which))
                attrs += f' marker-{which}="url(#{mid})"'
        return attrs

    @staticmethod
    def fill(sp_pr):
        sf = sp_pr.find(qn("a:solidFill"))
        if sf is None:
            return 'fill="none"'
        return f'fill="#{sf.find(qn("a:srgbClr")).get("val")}"'

    def sp(self, el, tf, connector):
        sp_pr = el.find(qn("p:spPr"))
        x, y, w, h, fh, fv = self.box(el, tf)
        stroke = self.stroke(sp_pr)
        geom = sp_pr.find(qn("a:prstGeom"))
        cust = sp_pr.find(qn("a:custGeom"))
        if connector:
            x1, x2 = (x + w, x) if fh else (x, x + w)
            y1, y2 = (y + h, y) if fv else (y, y + h)
            self.out.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" {stroke}/>')
        elif cust is not None:
            path = cust.find(".//" + qn("a:path"))
            pw, ph = int(path.get("w")) or 1, int(path.get("h")) or 1
            d = []
            for seg in path:
                pt = seg.find(qn("a:pt"))
                if pt is None:
                    continue
                cmd = "M" if seg.tag == qn("a:moveTo") else "L"
                d.append(f'{cmd}{x + int(pt.get("x")) * w / pw:.0f},{y + int(pt.get("y")) * h / ph:.0f}')
            self.out.append(f'<path d="{" ".join(d)}" {self.fill(sp_pr)} {stroke}/>')
        elif geom is not None:
            prst = geom.get("prst")
            if prst == "ellipse":
                self.out.append(f'<ellipse cx="{x + w / 2}" cy="{y + h / 2}" rx="{w / 2}" ry="{h / 2}" '
                                f'{self.fill(sp_pr)} {stroke}/>')
            elif prst in ("rect", "roundRect"):
                r = 0
                if prst == "roundRect":
                    gd = geom.find(".//" + qn("a:gd"))
                    adj = int(gd.get("fmla").split()[1]) / 100000 if gd is not None else 0.16667
                    r = adj * min(w, h)
                if el.find(qn("p:nvSpPr")).find(qn("p:cNvSpPr")).get("txBox") != "1":
                    self.out.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{r}" '
                                    f'{self.fill(sp_pr)} {stroke}/>')
        self.text(el, x, y, w, h)

    def text(self, el, x, y, w, h):
        tx = el.find(qn("p:txBody"))
        if tx is None:
            return
        body = tx.find(qn("a:bodyPr"))
        top = y + int(body.get("tIns", "45720"))
        left, right = x + int(body.get("lIns", "91440")), x + w - int(body.get("rIns", "91440"))
        lines = []
        for p in tx.findall(qn("a:p")):
            runs = p.findall(qn("a:r"))
            if not runs:
                continue
            ppr = p.find(qn("a:pPr"))
            algn = ppr.get("algn", "l") if ppr is not None else "l"
            rpr = runs[0].find(qn("a:rPr"))
            size = int(rpr.get("sz", "1800")) * 127  # 1/100 pt -> EMU
            bold = rpr.get("b") == "1"
            sf = rpr.find(qn("a:solidFill"))
            col = "#" + sf.find(qn("a:srgbClr")).get("val") if sf is not None else "#000000"
            s = "".join(r.find(qn("a:t")).text or "" for r in runs)
            lines.append((s, size, bold, col, algn))
        if body.get("anchor") == "ctr" and lines:
            top = y + (h - sum(l[1] * 1.2 for l in lines)) / 2
        for s, size, bold, col, algn in lines:
            base = top + ASCENT * size
            ax, anchor = {"ctr": ((left + right) / 2, "middle"), "r": (right, "end")}.get(algn, (left, "start"))
            weight = ' font-weight="bold"' if bold else ""
            self.out.append(f'<text x="{ax}" y="{base}" font-size="{size}" text-anchor="{anchor}" '
                            f'fill="{col}"{weight}>{escape(s)}</text>')
            top += size * 1.2

    def pic(self, el, tf):
        x, y, w, h, _, _ = self.box(el, tf)
        blip = el.find(".//" + qn("a:blip"))
        svg = blip.find(".//" + ASVG_BLIP)
        r_id, mime = (svg.get(REL_EMBED), "image/svg+xml") if svg is not None else (blip.get(REL_EMBED), "image/png")
        blob = self.slide.part.related_part(r_id).blob
        uri = f"data:{mime};base64,{base64.b64encode(blob).decode('ascii')}"
        self.out.append(f'<image href="{uri}" x="{x}" y="{y}" width="{w}" height="{h}"/>')

    def svg(self, sw, sh):
        defs = []
        for mid, c, which in sorted(self.markers):
            path = "M0,0 L8,3 L0,6 Z" if which == "end" else "M8,0 L0,3 L8,6 Z"
            ref = 8 if which == "end" else 0
            defs.append(f'<marker id="{mid}" markerWidth="10" markerHeight="10" refX="{ref}" refY="3" '
                        f'orient="auto"><path d="{path}" fill="#{c}"/></marker>')
        return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {sw} {sh}" '
                f'font-family="\'{self.font}\', sans-serif"><defs>{"".join(defs)}</defs>'
                f'<rect width="{sw}" height="{sh}" fill="#FFFFFF"/>{"".join(self.out)}</svg>')


def render(deck, stem):
    """Render every slide of `deck` to <stem>-1.png, <stem>-2.png, ...; return the paths."""
    prs = Presentation(str(deck))
    stem = Path(stem)
    font = preview_font()
    sw, sh = prs.slide_width, prs.slide_height
    outs = []
    for i, slide in enumerate(prs.slides, 1):
        r = Renderer(slide, font)
        r.walk(slide.shapes, lambda x, y, w, h: (x, y, w, h))
        out = stem.with_name(f"{stem.name}-{i}.png")
        cairosvg.svg2png(bytestring=r.svg(sw, sh).encode("utf-8"), write_to=str(out), output_width=2000)
        outs.append(out)
    return outs


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    for out in render(sys.argv[1], sys.argv[2]):
        print(out)


if __name__ == "__main__":
    main()
