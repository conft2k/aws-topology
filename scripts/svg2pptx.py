#!/usr/bin/env python3
"""Turn layout SVGs into a PowerPoint deck of native, editable shapes.

    uv run --with python-pptx --with cairosvg python scripts/svg2pptx.py \\
        draft.svg [more.svg ...] -o deck.pptx [--title "Slide title" ...]

Each SVG becomes one 16:9 slide. The diagram is scaled to fit the slide (below the
title, if one is given) and placed as one group, so it moves and resizes as a unit:

  rect      -> rectangle (square, or rounded when rx is set), exact fill, border
               color, width and dash pattern
  image     -> picture holding the official SVG icon (PowerPoint 2016+/365 draws the
               vector; older viewers use the embedded PNG fallback)
  text      -> text box with the same size, weight, color and anchor
  line      -> straight connector; marker-start / marker-end become arrowheads
  polyline  -> open freeform with the same arrowheads
  circle    -> oval (step badges, junction dots)
  g         -> its translate() is applied to the children

Icons are referenced exactly as in the SVG workflow: `bundle:set/name` hrefs (looked
up in assets/icons.json) or data: URIs. Nothing is written, and the script exits with
status 1, when an href doesn't resolve, points at an external file, names the
logo-less AWS Cloud badge, when a group container is rounded, or when the SVG uses an
element, transform, style/class or text stroke that this converter would otherwise
drop silently.
"""
import argparse
import base64
import io
import re
import sys
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import macos_cairo  # noqa: E402,F401  (before cairosvg: finds Homebrew cairo on Apple Silicon)
import cairosvg  # noqa: E402
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from fetch_icons import load_sets  # noqa: E402

SETS = load_sets()  # downloads the official AWS icons on first use
# An already-embedded SVG (data: URI from embed.py) maps back to its icon.
BY_DATA = {e["data"]: f"{setname}/{name}" for setname, icons in SETS.items()
           for name, e in icons.items()}

NS = "{http://www.w3.org/2000/svg}"
XLINK = "{http://www.w3.org/1999/xlink}href"
SVG_BLIP_EXT = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
WHITE = {"#fff", "#ffffff", "white"}
# The AWS Cloud container always carries the AWS logo; these are the logo-less outlines.
REPLACED = {"group/AWS-Cloud_32": "group/AWS-Cloud-logo_32",
            "group/AWS-Cloud_32_Dark": "group/AWS-Cloud-logo_32_Dark"}
DRAWN = {"svg", "g", "rect", "image", "text", "line", "polyline", "circle"}
SKIPPED = {"defs", "marker", "title", "desc"}  # markers are read through marker-start/-end

SLIDE_W, SLIDE_H = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in, 16:9
MARGIN = Emu(228600)                              # 0.25 in
TITLE_H = Emu(685800)                             # 0.75 in title band
# Baseline of a single text line, measured from the top of a zero-inset text box,
# as a fraction of the font size (Arial ascent 0.905, Malgun Gothic ~1.0).
ASCENT = 0.95
PNG_FALLBACK_PX = 256
MIN_PT = 10   # every label on the slide should be at least this size
FLOOR_PT = 8  # allowed only when the diagram cannot be tightened or split; never below


def num(v, default=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def color(v):
    """'#RRGGBB' / '#RGB' / 'white' -> RGBColor, 'none' or missing -> None."""
    v = (v or "none").strip().lower()
    if v == "none":
        return None
    if v in WHITE:
        return RGBColor(0xFF, 0xFF, 0xFF)
    if v in ("black",):
        return RGBColor(0, 0, 0)
    m = re.fullmatch(r"#([0-9a-f]{3}|[0-9a-f]{6})", v)
    if not m:
        raise ValueError(f"unsupported color {v!r} (use #RRGGBB)")
    h = m.group(1)
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    return RGBColor.from_string(h.upper())


def text_width(s, fs):
    """Same estimate as check_layout.py: 0.55 em Latin, 1 em CJK."""
    w = 0.0
    for c in s:
        wide = "ᄀ" <= c <= "ᇿ" or "　" <= c <= "鿿" or "가" <= c <= "힣"
        w += fs * (1.0 if wide else 0.55)
    return w


def translate(el):
    t = el.get("transform", "").strip()
    if not t:
        return 0.0, 0.0
    m = re.fullmatch(r"translate\(\s*([-\d.]+)(?:[ ,]+([-\d.]+))?\s*\)", t)
    if not m:
        raise ValueError(f"unsupported transform {t!r} (only translate(x,y))")
    return num(m.group(1)), num(m.group(2))


def points(s):
    vals = [num(v) for v in re.split(r"[ ,]+", (s or "").strip()) if v]
    return [(vals[i], vals[i + 1]) for i in range(0, len(vals) - 1, 2)]


def has_marker(el, which):
    return bool(re.match(r"url\(#", el.get(which, "")))


def resolve_icon(href):
    """href -> (cache key, svg text or None, png bytes or None), or raise ValueError."""
    if href.startswith("bundle:"):
        ref = href[len("bundle:"):]
        if ref in REPLACED:
            raise ValueError(f"bundle:{ref} has no AWS logo -> use bundle:{REPLACED[ref]}")
        setname, _, name = ref.partition("/")
        entry = SETS.get(setname, {}).get(name)
        if entry is None:
            raise ValueError(f"bundle:{ref} is not in the bundle (check with find_icon.py)")
        if entry["type"] == "svg":
            return ref, entry["data"], None
        return ref, None, base64.b64decode(entry["data"])
    m = re.match(r"data:image/(svg\+xml|png);base64,(.*)", href, re.S)
    if m:
        raw = base64.b64decode(m.group(2))
        if m.group(1) == "svg+xml":
            svg = raw.decode("utf-8")
            return BY_DATA.get(svg, "data/icon-%d" % hash(raw)), svg, None
        return BY_DATA.get(m.group(2), "data/icon-%d" % hash(raw)), None, raw
    raise ValueError(f"image href {href[:60]!r} is external (use bundle:set/name)")


def collect(svg_root):
    """Flatten the SVG into drawing items in document order, plus every problem found."""
    items, problems = [], []
    vb = [num(v) for v in svg_root.get("viewBox", "").replace(",", " ").split()]
    if len(vb) != 4:
        vb = [0, 0, num(svg_root.get("width")), num(svg_root.get("height"))]
    vx, vy, vw, vh = vb
    if vw <= 0 or vh <= 0:
        problems.append("the root <svg> needs a viewBox (or width and height)")
        return (vx, vy, vw, vh), items, problems
    badges = set()

    def walk(el, ox, oy):
        tag = el.tag.replace(NS, "")
        if tag in SKIPPED:
            return
        if tag not in DRAWN:
            problems.append(f"<{tag}> is not supported (draw it with rect/line/polyline/circle/text/image)")
            return
        try:
            dx, dy = translate(el)
        except ValueError as e:
            problems.append(f"<{tag}>: {e}")
            return
        ox, oy = ox + dx, oy + dy
        g = el.get
        if g("style") or g("class"):
            problems.append(f"<{tag}> uses style=/class= — write presentation attributes "
                            "(fill=, stroke=, font-size=, ...) instead")
        try:
            if tag == "rect":
                x, y, w, h = num(g("x")) + ox, num(g("y")) + oy, num(g("width")), num(g("height"))
                items.append({"k": "rect", "box": (x, y, w, h), "fill": color(g("fill", "#000")),
                              "stroke": color(g("stroke")), "sw": num(g("stroke-width"), 1),
                              "dash": g("stroke-dasharray"), "r": num(g("rx") or g("ry"))})
            elif tag == "image":
                href = g("href") or g(XLINK) or ""
                x, y = num(g("x")) + ox, num(g("y")) + oy
                if href.startswith("bundle:group/"):
                    badges.add((round(x, 2), round(y, 2)))
                key, svg, png = resolve_icon(href)
                items.append({"k": "image", "box": (x, y, num(g("width")), num(g("height"))),
                              "key": key, "svg": svg, "png": png,
                              "name": key.partition("/")[2]})
            elif tag == "text":
                s = "".join(el.itertext()).strip()
                if color(g("stroke")) is not None:
                    problems.append(f"<text> {s!r} has a stroke (text halo) — PowerPoint text has no "
                                    "stroke; put a white mask rect behind the label instead")
                elif s:
                    items.append({"k": "text", "s": s, "x": num(g("x")) + ox, "y": num(g("y")) + oy,
                                  "fs": num(g("font-size"), 16), "fill": color(g("fill", "#000")),
                                  "bold": g("font-weight", "normal") in ("bold", "bolder", "600", "700", "800", "900"),
                                  "anchor": g("text-anchor", "start")})
                return  # tspans are part of the text run
            elif tag in ("line", "polyline"):
                if tag == "line":
                    pts = [(num(g("x1")), num(g("y1"))), (num(g("x2")), num(g("y2")))]
                else:
                    pts = points(g("points"))
                    if color(g("fill")) is not None:
                        problems.append("<polyline> with a fill is not supported (use fill=\"none\")")
                items.append({"k": "path", "pts": [(px + ox, py + oy) for px, py in pts],
                              "stroke": color(g("stroke", "#000")), "sw": num(g("stroke-width"), 1),
                              "dash": g("stroke-dasharray"),
                              "head": has_marker(el, "marker-start"), "tail": has_marker(el, "marker-end")})
            elif tag == "circle":
                cx, cy, r = num(g("cx")) + ox, num(g("cy")) + oy, num(g("r"))
                items.append({"k": "oval", "box": (cx - r, cy - r, 2 * r, 2 * r), "fill": color(g("fill", "#000")),
                              "stroke": color(g("stroke")), "sw": num(g("stroke-width"), 1), "dash": None})
        except ValueError as e:
            problems.append(f"<{tag}>: {e}")
            return
        for child in el:
            walk(child, ox, oy)

    walk(svg_root, -vx, -vy)
    for it in items:
        if it["k"] == "rect" and it["r"] > 0 and (round(it["box"][0], 2), round(it["box"][1], 2)) in badges:
            problems.append(f"rounded group container at ({it['box'][0]:g},{it['box'][1]:g}) — "
                            "group containers are square-cornered, drop rx/ry")
    # The full-canvas white background is the slide's own background.
    items = [it for it in items if not (it["k"] == "rect" and it["box"][2] >= vw and it["box"][3] >= vh
                                        and it["stroke"] is None)]
    return (0.0, 0.0, vw, vh), items, problems


class Deck:
    def __init__(self, font, ea_font):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_W, SLIDE_H
        # python-pptx's template carries its own author and 2013 dates; reset them.
        props, now = self.prs.core_properties, datetime.now(timezone.utc).replace(tzinfo=None)
        props.author = props.last_modified_by = ""
        props.comments = "Generated by skill-aws-arch-diagram"
        props.created = props.modified = now
        props.revision = 1
        self.font, self.ea_font = font, ea_font
        self.svg_parts = {}   # icon key -> SVG part, shared by every slide
        self.png_cache = {}   # icon key -> PNG fallback bytes

    # ---- placement -------------------------------------------------------
    def add_slide(self, view, items, title):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])  # blank
        top = MARGIN
        if title:
            self._title(slide, title)
            top = MARGIN + TITLE_H
        _, _, vw, vh = view
        avail_w, avail_h = SLIDE_W - 2 * MARGIN, SLIDE_H - top - MARGIN
        self.k = min(avail_w / vw, avail_h / vh)          # EMU per SVG unit
        self.ox = MARGIN + (avail_w - vw * self.k) / 2
        self.oy = top + (avail_h - vh * self.k) / 2
        group = slide.shapes.add_group_shape()
        group.name = "Architecture diagram"
        for it in items:
            getattr(self, "_" + it["k"])(slide, group.shapes, it)
        return slide

    def X(self, v):
        return Emu(int(round(self.ox + v * self.k)))

    def Y(self, v):
        return Emu(int(round(self.oy + v * self.k)))

    def L(self, v):
        return Emu(max(1, int(round(v * self.k))))

    def pt(self, px):
        return px * self.k / 12700

    # ---- line / fill helpers --------------------------------------------
    def _ln(self, sp_pr, stroke, sw, dash, head=False, tail=False):
        for old in sp_pr.findall(qn("a:ln")):
            sp_pr.remove(old)
        ln = etree.SubElement(sp_pr, qn("a:ln"))
        if stroke is None:
            etree.SubElement(ln, qn("a:noFill"))
            return
        ln.set("w", str(int(round(self.pt(sw) * 12700))))
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", str(stroke))
        parts = [num(v) for v in re.split(r"[ ,]+", (dash or "").strip()) if v and v != "none"]
        if parts:
            if len(parts) % 2:
                parts *= 2
            cust = etree.SubElement(ln, qn("a:custDash"))
            for d, s in zip(parts[::2], parts[1::2]):
                ds = etree.SubElement(cust, qn("a:ds"))
                ds.set("d", str(int(round(d / sw * 100000))))
                ds.set("sp", str(int(round(s / sw * 100000))))
        else:
            etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        etree.SubElement(ln, qn("a:miter")).set("lim", "800000")
        for tag, on in (("a:headEnd", head), ("a:tailEnd", tail)):
            if on:
                end = etree.SubElement(ln, qn(tag))
                end.set("type", "triangle")
                end.set("w", "med")
                end.set("len", "med")

    @staticmethod
    def _fill(shape, rgb):
        if rgb is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb

    @staticmethod
    def _plain(shape):
        """Drop the theme style reference so no theme line, fill or shadow leaks in."""
        style = shape._element.find(qn("p:style"))
        if style is not None:
            shape._element.remove(style)

    # ---- items -----------------------------------------------------------
    def _rect(self, slide, shapes, it):
        x, y, w, h = it["box"]
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if it.get("r", 0) > 0 else MSO_SHAPE.RECTANGLE
        shape = shapes.add_shape(kind, self.X(x), self.Y(y), self.L(w), self.L(h))
        self._plain(shape)
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            shape.adjustments[0] = min(0.5, it["r"] / min(w, h))
        self._fill(shape, it["fill"])
        self._ln(shape._element.spPr, it["stroke"], it["sw"], it["dash"])
        return shape

    def _oval(self, slide, shapes, it):
        x, y, w, h = it["box"]
        shape = shapes.add_shape(MSO_SHAPE.OVAL, self.X(x), self.Y(y), self.L(w), self.L(h))
        self._plain(shape)
        self._fill(shape, it["fill"])
        self._ln(shape._element.spPr, it["stroke"], it["sw"], it["dash"])

    def _path(self, slide, shapes, it):
        pts = it["pts"]
        if len(pts) == 2:
            (x1, y1), (x2, y2) = pts
            shape = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self.X(x1), self.Y(y1), self.X(x2), self.Y(y2))
        else:
            ff = shapes.build_freeform(self.X(pts[0][0]), self.Y(pts[0][1]), scale=1.0)
            ff.add_line_segments([(self.X(px), self.Y(py)) for px, py in pts[1:]], close=False)
            shape = ff.convert_to_shape()
            shape.fill.background()
        self._plain(shape)
        self._ln(shape._element.spPr, it["stroke"], it["sw"], it["dash"], it["head"], it["tail"])

    def _text(self, slide, shapes, it):
        fs, s = it["fs"], it["s"]
        w = text_width(s, fs) * 1.25 + fs  # slack so a wider real font doesn't wrap or clip
        x = it["x"] - {"middle": w / 2, "end": w}.get(it["anchor"], 0)
        box = shapes.add_textbox(self.X(x), self.Y(it["y"] - ASCENT * fs), self.L(w), self.L(fs * 1.3))
        tf = box.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p = tf.paragraphs[0]
        p.alignment = {"middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}.get(it["anchor"], PP_ALIGN.LEFT)
        p.line_spacing = 1.0
        run = p.add_run()
        run.text = s
        self._font(run, Pt(self.pt(fs)), it["bold"], it["fill"])

    def _font(self, run, size, bold, rgb):
        f = run.font
        f.size, f.bold, f.name = size, bold, self.font
        if rgb is not None:
            f.color.rgb = rgb
        r_pr = run._r.get_or_add_rPr()
        r_pr.set("lang", "ko-KR")
        for tag in ("a:ea", "a:cs"):
            el = r_pr.find(qn(tag))
            if el is None:
                el = etree.SubElement(r_pr, qn(tag))
            el.set("typeface", self.ea_font)

    def _image(self, slide, shapes, it):
        x, y, w, h = it["box"]
        key = it["key"]
        if key not in self.png_cache:
            self.png_cache[key] = it["png"] or cairosvg.svg2png(
                bytestring=it["svg"].encode("utf-8"), output_width=PNG_FALLBACK_PX,
                output_height=int(round(PNG_FALLBACK_PX * h / w)) if w else None)
        pic = shapes.add_picture(io.BytesIO(self.png_cache[key]), self.X(x), self.Y(y), self.L(w), self.L(h))
        pic.name = it["name"]
        pic._element.nvPicPr.cNvPr.set("descr", it["name"].replace("-", " "))
        if it["svg"] is None:
            return
        part = self.svg_parts.get(key)
        if part is None:
            package = self.prs.part.package
            part = Part(package.next_image_partname("svg"), "image/svg+xml", package,
                        it["svg"].encode("utf-8"))
            self.svg_parts[key] = part
        r_id = slide.part.relate_to(part, RT.IMAGE)
        blip = pic._element.blipFill.find(qn("a:blip"))
        ext_lst = etree.SubElement(blip, qn("a:extLst"))
        ext = etree.SubElement(ext_lst, qn("a:ext"))
        ext.set("uri", SVG_BLIP_EXT)
        svg_blip = etree.SubElement(ext, f"{{{ASVG}}}svgBlip", nsmap={"asvg": ASVG})
        svg_blip.set(qn("r:embed"), r_id)

    def _title(self, slide, title):
        box = slide.shapes.add_textbox(MARGIN, MARGIN, SLIDE_W - 2 * MARGIN, TITLE_H)
        box.name = "Title"
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        run = tf.paragraphs[0].add_run()
        run.text = title
        self._font(run, Pt(24), True, RGBColor(0x23, 0x2F, 0x3E))


def build(svgs, titles, out, font="Arial", ea_font="Malgun Gothic"):
    """Convert layout SVGs into a deck at `out`; return False (writing nothing) on a bad SVG."""
    if len(titles) > len(svgs):
        print("more --title values than SVGs")
        return False
    parsed, failed = [], False
    for path in svgs:
        view, items, problems = collect(ET.parse(path).getroot())
        if problems:
            failed = True
            print(f"{path}:")
            for p in dict.fromkeys(problems):
                print(f"  {p}")
        parsed.append((view, items))
    if failed:
        if out.exists():
            out.unlink()  # never leave a stale deck behind
        print(f"Nothing written to {out}; fix the SVG(s) above and re-run.")
        return False

    deck = Deck(font, ea_font)
    deck.prs.core_properties.title = titles[0] if titles else ""
    for i, (view, items) in enumerate(parsed):
        deck.add_slide(view, items, titles[i] if i < len(titles) else None)
        sizes = [deck.pt(it["fs"]) for it in items if it["k"] == "text"]
        if sizes and min(sizes) < FLOOR_PT:
            print(f"WARNING: slide {i + 1} ({svgs[i].name}): smallest text is {min(sizes):.1f}pt, under the "
                  f"{FLOOR_PT}pt floor — not deliverable; tighten the canvas or split the diagram")
        elif sizes and min(sizes) < MIN_PT:
            print(f"note: slide {i + 1} ({svgs[i].name}): smallest text is {min(sizes):.1f}pt (< {MIN_PT}pt) — "
                  f"tighten or split it; keep {FLOOR_PT}-{MIN_PT}pt only if unavoidable and tell the user")
    deck.prs.save(out)
    icons = sum(1 for _, items in parsed for it in items if it["k"] == "image")
    print(f"Wrote {len(parsed)} slide(s), {icons} icon(s) ({len(deck.svg_parts)} distinct SVGs) -> {out}")
    return True


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("svgs", nargs="+", type=Path, help="layout SVGs, one slide each, in order")
    ap.add_argument("-o", "--out", required=True, type=Path, help="output .pptx")
    ap.add_argument("--title", action="append", default=[],
                    help="slide title; repeat once per slide in order (omit for no title band)")
    ap.add_argument("--font", default="Arial", help="Latin typeface (default Arial)")
    ap.add_argument("--ea-font", default="Malgun Gothic",
                    help="Korean/East Asian typeface (default Malgun Gothic, i.e. 맑은 고딕)")
    a = ap.parse_args()
    if not build(a.svgs, a.title, a.out, a.font, a.ea_font):
        sys.exit(1)


if __name__ == "__main__":
    main()
